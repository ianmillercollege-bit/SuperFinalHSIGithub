"""BACKEND_CONTRACT.md v1.4.1 section 7c, "Seed at scale" (test_seed_scale)."""

import json
import re
import subprocess
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select

from db import Brand, ComparisonFact, DailyMetric, Incident, Product, SessionLocal, User
from seed_loader import rebuild_database
from services.checker import Catalog, check, extract_claims
from services.passwords import verify_password

BACKEND = Path(__file__).resolve().parents[1]
DATA = BACKEND / "seed" / "data"
CATALOG = DATA / "catalog"
SOURCE = BACKEND / "seed" / "source" / "greek_god_tech_companies.xlsx"
REAL_BRANDS = {"nike", "hermes", "aura", "nyx", "eos", "iris", "atlas", "apollo", "kratos", "apple", "samsung",
               "sony", "dell", "lenovo", "asus", "acer", "microsoft", "google", "bose", "logitech", "razer", "intel"}


def rows(name: str) -> list[dict]:
    return next(v for v in json.loads((CATALOG / f"{name}.json").read_text(encoding="utf-8")).values())


@pytest.fixture
def seeded(client, monkeypatch):
    monkeypatch.setenv("SEED_DIR", str(DATA))
    rebuild_database()
    return client


def test_seed_scale(seeded):
    with SessionLocal() as db:
        brands = db.scalars(select(Brand)).all()
        assert len(brands) == 153
        sheet_products = db.scalar(select(func.count()).select_from(Product).where(Product.product_id.like("prod_%-%")))
        assert sheet_products == 1500
        assert db.scalar(select(func.count()).select_from(Product)) == 1512  # plus the 12 originals

        # Every brand has at least one admin, a trend and incidents.
        admins = {u.brand_id for u in db.scalars(select(User)).all() if u.brand_id}
        trends = set(db.scalars(select(DailyMetric.brand_id).distinct()).all())
        with_incidents = set(db.scalars(select(Incident.brand_id).distinct()).all())
        for b in brands:
            assert b.brand_id in admins, b.name
            assert b.brand_id in trends, b.name
            assert b.brand_id in with_incidents, b.name

        # No real brand names in company names.
        for b in brands:
            words = set(re.findall(r"[a-z]+", b.name.lower()))
            assert not words & REAL_BRANDS, b.name


def test_rebuild_is_under_ten_seconds(monkeypatch):
    monkeypatch.setenv("SEED_DIR", str(DATA))
    start = time.perf_counter()
    rebuild_database()
    assert time.perf_counter() - start < 10


def test_every_company_has_full_dashboard_data():
    """Contract: 8+ answers, 20+ claims, 4+ incidents covering the four statuses, for every company (the
    rebuild is well under 10 s, so no company needs the reduced fallback)."""
    answers, claims, incidents = Counter(), Counter(), defaultdict(list)
    claim_brand = {}
    for a in rows("answers"):
        answers[a["brandId"]] += 1
        claim_brand[a["answerId"]] = a["brandId"]
    for c in rows("claims"):
        claims[claim_brand[c["answerId"]]] += 1
    for i in rows("incidents"):
        incidents[i["brandId"]].append(i["status"])
    companies = [b["brandId"] for b in rows("brands")]
    assert len(companies) == 150
    for brand_id in companies:
        assert answers[brand_id] >= 8 and claims[brand_id] >= 20 and len(incidents[brand_id]) >= 4, brand_id
        assert {"auto_fixed", "pending_approval", "approved", "rejected"} <= set(incidents[brand_id]), brand_id
    assert sum("escalated" in statuses for statuses in incidents.values()) >= 5


def test_products_follow_the_sheet(seeded):
    by_category = Counter(p["category"] for p in rows("products"))
    assert by_category == {"headphones": 250, "laptops": 250, "phones_tablets": 500, "computer_hardware": 500}
    headphones = seeded.get("/api/v1/products", params={"category": "headphones"}).json()["products"]
    assert len(headphones) == 250 and {p["category"] for p in headphones} == {"headphones"}
    sample = headphones[0]
    assert sample["subcategory"] in {"Headphones", "Earbuds", "Headset"}
    assert set(sample["specs"]) <= {"processor", "graphics", "displayType", "resolution", "ports", "operatingSystem",
                                    "batteryHours", "weightG", "warranty", "certifications", "useCaseTags",
                                    "otherNames"}
    assert isinstance(sample["specs"]["certifications"], list) and isinstance(sample["specs"]["otherNames"], list)
    laptop = seeded.get("/api/v1/products", params={"category": "laptops", "brandId": "brand_004"})
    assert laptop.status_code == 200
    sheet_laptop = next(p for p in seeded.get("/api/v1/products", params={"category": "laptops"}).json()["products"]
                        if "-" in p["productId"])
    for key in ("ramGb", "storageGb", "screenInches", "batteryHours", "weightLb", "touchscreen", "processor"):
        assert key in sheet_laptop["specs"], key
    # Originals keep exactly their v1.3 specs.
    aero = next(p for p in seeded.get("/api/v1/products").json()["products"] if p["productId"] == "prod_001")
    assert set(aero["specs"]) == {"ramGb", "storageGb", "screenInches", "batteryHours", "weightLb", "touchscreen"}
    assert seeded.get("/api/v1/products", params={"brandId": "brand_999"}).status_code == 404


def test_verified_comparisons_are_facts(seeded):
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(ComparisonFact)) >= 2900
        fact = db.scalars(select(ComparisonFact).where(ComparisonFact.attribute == "price")).first()
        a, b = db.get(Product, fact.product_id), db.get(Product, fact.other_product_id)
        catalog = Catalog(db)
        supported = [check(c, catalog) for c in extract_claims(f"The {a.name} is cheaper than the {b.name}.", catalog)]
        assert [(r.status, r.rule_id, r.fact_id) for r in supported] == [("correct", None, fact.fact_id)]
        unsupported = [check(c, catalog) for c in extract_claims(f"The {b.name} is cheaper than the {a.name}.", catalog)]
        assert [(r.status, r.rule_id) for r in unsupported] == [("incorrect", "UNFAIR_COMPARISON")]


def test_demo_users_and_passwords(seeded):
    with SessionLocal() as db:
        users = db.scalars(select(User)).all()
        assert len(users) >= 158
        staff = {u.username: u for u in users if u.role == "CIRQO Staff"}
        assert set(staff) == {"grace.kim@cirqo.example", "dev.patel@cirqo.example"}
        assert all(u.brand_id is None for u in staff.values())
        maria = next(u for u in users if u.username == "maria.lopez@kestrel.example")
        assert (maria.brand_id, maria.role) == ("brand_001", "Brand Data Owner")
        for u in users:
            assert u.password_hash.startswith("pbkdf2_sha256$") and "cirqo-demo" not in u.password_hash
        assert verify_password("cirqo-demo", users[0].password_hash)
        assert not verify_password("wrong", users[0].password_hash)
        sheet_admin = db.scalars(select(User).where(User.brand_id == "brand_004")).one()
        assert sheet_admin.username == "admin@morpheusaudio.example" and sheet_admin.role == "Brand Data Owner"


def test_committed_sheet_has_no_passwords():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["Login Credentials"]
    header = next(r for r in ws.iter_rows(values_only=True) if r and "Login Email" in r)
    column = next(i for i, h in enumerate(header) if h and h.startswith("Temporary Password"))
    values = [r[column] for r in ws.iter_rows(values_only=True) if r and r[0] and str(r[0]).startswith("CO-")]
    assert len(values) == 150 and not any(values)


def test_existing_brands_are_unchanged(seeded):
    names = {b["brandId"]: b["brandName"] for b in seeded.get("/api/v1/auth/demo-accounts").json()["accounts"]}
    assert {k: names[k] for k in ("brand_001", "brand_002", "brand_003")} ==         {"brand_001": "Kestrel", "brand_002": "Arcton", "brand_003": "Novex"}
    competitors = seeded.get("/api/v1/visibility/summary").json()["competitors"]
    assert [c["brandName"] for c in competitors] == ["Arcton", "Novex"]
    # A laptop question never returns earbuds; a headphone question returns headphones.
    laptop = seeded.post("/api/v1/connector/query", json={"question": "Best laptop under $500 for school?",
                                                          "assistantId": "ast_01"}).json()
    products = {p["productId"]: p for p in seeded.get("/api/v1/products").json()["products"]}
    assert products[laptop["recommendation"]["productId"]]["category"] == "laptops"
    phones = seeded.post("/api/v1/connector/query", json={"question": "Good headphones for the gym under $150?",
                                                          "assistantId": "ast_01"}).json()
    assert products[phones["recommendation"]["productId"]]["category"] == "headphones"
    assert all(c["status"] == "correct" for c in phones["claims"])


def test_catalog_files_are_deterministic_and_up_to_date(tmp_path):
    subprocess.run([sys.executable, str(BACKEND / "seed" / "generate.py"), "--out", str(tmp_path)],
                   check=True, capture_output=True)
    for f in CATALOG.glob("*.json"):
        fresh = (tmp_path / "catalog" / f.name).read_bytes().replace(b"\r\n", b"\n")
        assert fresh == f.read_bytes().replace(b"\r\n", b"\n"), f"catalog/{f.name} is stale: rerun generate.py"
